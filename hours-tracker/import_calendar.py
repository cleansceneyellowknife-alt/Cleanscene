"""Pre-fill the Hours Log from Google Calendar job invites.

Every job on the CleanScene calendar invites the crew by email. Each invited crew member
on an event becomes one Hours Log row (Date, Employee, Client / Site, Start, End, Break 0)
with "From calendar: <event title>" in column J. Ali then only corrects shifts that ran
long, short or were cancelled.

Usage:
    python3 import_calendar.py <calendar_events.json> <from YYYY-MM-DD> <to YYYY-MM-DD> [workbook.xlsx]

<calendar_events.json> is a Google Calendar events export ({"events": [...]}, as returned by
the Calendar API / connector). Rows already imported (same date, employee, site, start) are skipped.
The calendar export holds client door codes: keep it out of the repo.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
FIRST, LAST = 2, 2001

# Crew calendar emails -> Staff tab names.
CREW = {
    "chalijattbabu@gmail.com": "Ali",
    "villanuevaermalyn29@gmail.com": "Ermalyn Villanueva",
    "alijan0987as@gmail.com": "Ali",  # Ali's second account
}

# (employee, site) pairs never imported: Ermalyn is not tracked on NACC jobs (Adam, 2026-10-12).
SKIP = {("Ermalyn Villanueva", "NACC")}

# First matching keyword (in event title or location, case-insensitive) -> Sites tab name.
SITE_KEYWORDS = [
    ("allain", "École Allain St-Cyr, CSFTNO"),
    ("day care", "La Taniere, CSFTNO"),  # the daycare at 5102 53 St is La Taniere
    ("melissa", "Melissa Kruger"),
    ("athena", "Athena Wells"),
    ("co operators", "The Co-operators c/o Nathalie Gauthier Dore"),
    ("321d old airport", "The Co-operators c/o Nathalie Gauthier Dore"),
    ("beverly", "Beverly Speight"),
    ("amy", "Amy Burt"),
    ("house clean - ben", "Ben Fraser"),
    ("4007 school draw", "5950 N.W.T. LIMITED"),
    ("carol", "Carol Pinto (Jamie Lavers)"),
    ("hospital", "Stanton Territorial Hospital"),
    ("shawn", "Shawn McCann"),
    ("dechinta", "Dechinta Centre for Research and Learning"),
    ("natasha", "Natasha Sanders"),
    ("9 nahanni", "Custom Electric Ltd."),
    ("susan", "Susan Eaton"),
    ("keith", "Keith MacNeil"),
    ("anita", "Anita Griffore"),
    ("stefanie", "Stefanie Pandke"),
    ("shelter", "Housing NWT Shelter / Sobering Centre"),
    ("nacc", "NACC"),
    ("stuart york", "Stuart York"),
    ("marks", "CT REIT - Mark's Yellowknife"),
    ("boston pizza", "Boston Pizza Yellowknife"),
    ("212 woolgar", "Danielle Desmarais"),
]


def site_for(event):
    hay = f"{event.get('summary', '')} {event.get('location') or ''}".lower()
    for key, site in SITE_KEYWORDS:
        if key in hay:
            return site
    return event.get("summary", "").strip()  # unknown: the Check column flags it


def rows_from(events, start, end):
    for ev in events:
        s, e = ev.get("start", {}).get("dateTime"), ev.get("end", {}).get("dateTime")
        if not s or not e:
            continue
        s, e = datetime.fromisoformat(s), datetime.fromisoformat(e)
        if not (start <= s.date() <= end):
            continue
        for att in ev.get("attendees") or []:
            name = CREW.get((att.get("email") or "").lower())
            if not name or att.get("responseStatus") == "declined":
                continue
            if (name, site_for(ev)) in SKIP:
                continue
            yield (s.date(), name, site_for(ev), s.time().replace(tzinfo=None),
                   e.time().replace(tzinfo=None), f"From calendar: {ev.get('summary', '').strip()}")


def main():
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    events = json.loads(Path(sys.argv[1]).read_text())["events"]
    start, end = date.fromisoformat(sys.argv[2]), date.fromisoformat(sys.argv[3])
    book = Path(sys.argv[4]) if len(sys.argv) > 4 else HERE / "CleanScene_Hours_Tracker.xlsx"
    wb = load_workbook(book)
    ws = wb["Hours Log"]

    def key(d, emp, site, st):
        return (d.isoformat() if hasattr(d, "isoformat") else str(d), emp, site, str(st)[:5])

    existing, next_row = set(), None
    for r in range(FIRST, LAST + 1):
        a = ws.cell(r, 1).value
        if a in (None, ""):
            next_row = next_row or r
            continue
        d = a.date() if isinstance(a, datetime) else a
        existing.add(key(d, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value))

    added = skipped = 0
    for d, emp, site, st, en, note in sorted(rows_from(events, start, end)):
        if key(d, emp, site, st) in existing:
            skipped += 1
            continue
        while next_row <= LAST and ws.cell(next_row, 1).value not in (None, ""):
            next_row += 1
        if next_row > LAST:
            sys.exit("Hours Log is full (row 2001).")
        for col, val in zip((1, 2, 3, 4, 5, 6, 10), (d, emp, site, st, en, 0, note)):
            ws.cell(next_row, col, val)
        existing.add(key(d, emp, site, st))
        next_row += 1
        added += 1
    wb.save(book)
    print(f"Added {added} shifts, skipped {skipped} already in the log -> {book}")


if __name__ == "__main__":
    main()
