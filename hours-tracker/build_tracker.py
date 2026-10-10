"""Build the CleanScene Hours Tracker workbook.

Usage:  python3 build_tracker.py            -> writes CleanScene_Hours_Tracker.xlsx
Then:   upload to Google Drive, open with Google Sheets, File > Save as Google Sheets,
        and share it with Adam and the manager only.

Formulas are Google Sheets compatible (SUMIFS/COUNTIFS/INDEX/MATCH only).
Staff = crew invited to jobs on the Google Calendar; Ali $28/h, floaters $25/h. Sites come from sites.json (pulled from Jobber).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT = HERE / "CleanScene_Hours_Tracker.xlsx"

LOG_FIRST, LOG_LAST = 2, 2001
STAFF_LAST = 101   # Staff!A2:A101 -> 100 people
SITES_LAST = 501   # Sites!A2:A501 -> 500 sites

NAVY = "1F3A5F"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")   # cells a person types into
FLAG_FILL = PatternFill("solid", fgColor="F8D7DA")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
BOLD = Font(bold=True)

DATE_FMT = "yyyy-mm-dd ddd"
TIME_FMT = "h:mm AM/PM"

LOG = "'Hours Log'"
RNG = lambda col: f"{LOG}!${col}${LOG_FIRST}:${col}${LOG_LAST}"
STAFF_NAMES = f"Staff!$A$2:$A${STAFF_LAST}"
STAFF_RATES = f"Staff!$D$2:$D${STAFF_LAST}"
SITE_NAMES = f"Sites!$A$2:$A${SITES_LAST}"
SITE_REGIONS = f"Sites!$B$2:$B${SITES_LAST}"

PLACEHOLDER_STAFF = [
    # name, region, role, hourly rate, active
    ("Ali", "Yellowknife", "Manager", 28.00, "Y"),
    ("Ermalyn Villanueva", "Yellowknife", "Floater", 25.00, "Y"),
    ("alijan0987as (confirm name)", "Yellowknife", "Floater", 25.00, "Y"),
]


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=row, column=i).column_letter].width = w


def build_start_here(ws):
    ws.column_dimensions["A"].width = 110
    lines = [
        ("CleanScene Hours Tracker", TITLE_FONT),
        ("Private: only Adam and the manager have access. Do not share this file.", BOLD),
        ("", None),
        ("MANAGER: logging a shift (2 minutes, end of each day)", BOLD),
        ("1. Open the 'Log a Shift' form (bookmark it on your phone), or go to the Hours Log tab.", None),
        ("2. One row = one person at one site for one shift.", None),
        ("3. Pick Employee and Client / Site from the dropdowns. Type Date, Start, End, Break minutes.", None),
        ("4. Add the Jobber visit # or a short note in column J (optional, helps the cross-check).", None),
        ("5. Grey columns (Hours, Region, Week of, Check) fill themselves. Never type in them.", None),
        ("6. If Check is not OK, fix the row the same day.", None),
        ("", None),
        ("RULES", BOLD),
        ("- Typing directly: use the first empty row. Never add rows after row 2001, never delete rows. Clear the cells instead.", None),
        ("- New employee: add to Staff tab first. New client: add to Sites tab first (same name as in Jobber).", None),
        ("- Overnight shifts work: an End earlier than Start is treated as the next day.", None),
        ("", None),
        ("ADAM: reviewing", BOLD),
        ("- Pay Period tab: hours, shifts, gross pay and flags per person for the current period. "
         "Pay is monthly: set 'Months back' to 1 to run last month's payroll.", None),
        ("- Weekly tab: hours per person per day, and the site cross-check against Jobber visits.", None),
        ("- Every Monday: open Weekly, set 'Weeks back' to 1, enter Jobber visits/hours for last week, clear every flag.", None),
        ("", None),
        ("CHECK column meanings", BOLD),
        ("OK = good.  Missing info = a required cell is blank.  Unknown employee / Unknown site = not on the Staff or Sites tab.", None),
        ("Check times = End is the same as Start or break is longer than the shift.  Over 12h = confirm it is real.", None),
        ("Duplicate? = same person, same site, same date twice.  Future date = date is after today.", None),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=r, column=1, value=text)
        if font:
            c.font = font


def build_hours_log(ws):
    header(ws, 1, ["Date", "Employee", "Client / Site", "Start", "End", "Break (min)", "Hours",
                   "Region", "Week of (Mon)", "Jobber visit # / note", "Check", "Visit helper"],
           [16, 22, 34, 11, 11, 11, 8, 18, 15, 26, 18, 10])
    ws.freeze_panes = "A2"
    grey = PatternFill("solid", fgColor="EEEEEE")
    for r in range(LOG_FIRST, LOG_LAST + 1):
        ws[f"G{r}"] = f'=IF(OR(A{r}="",D{r}="",E{r}=""),"",ROUND(MOD(E{r}-D{r},1)*24-N(F{r})/60,2))'
        ws[f"H{r}"] = f'=IF(C{r}="","",IFERROR(INDEX({SITE_REGIONS},MATCH(C{r},{SITE_NAMES},0)),"?"))'
        ws[f"I{r}"] = f'=IF(A{r}="","",A{r}-WEEKDAY(A{r},3))'
        ws[f"L{r}"] = f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm-dd")&"|"&B{r}&"|"&C{r})'
        ws[f"K{r}"] = (
            f'=IF(A{r}="","",'
            f'IF(OR(B{r}="",C{r}="",D{r}="",E{r}=""),"Missing info",'
            f'IF(ISNA(MATCH(B{r},{STAFF_NAMES},0)),"Unknown employee",'
            f'IF(H{r}="?","Unknown site",'
            f'IF(G{r}<=0,"Check times",'
            f'IF(G{r}>12,"Over 12h",'
            f'IF(COUNTIFS($L${LOG_FIRST}:$L${LOG_LAST},L{r})>1,"Duplicate?",'
            f'IF(A{r}>TODAY(),"Future date","OK"))))))))'
        )
        ws[f"A{r}"].number_format = DATE_FMT
        ws[f"I{r}"].number_format = "yyyy-mm-dd"
        ws[f"D{r}"].number_format = TIME_FMT
        ws[f"E{r}"].number_format = TIME_FMT
        ws[f"G{r}"].number_format = "0.00"
        for col in "GHIK":
            ws[f"{col}{r}"].fill = grey
    ws.column_dimensions["L"].hidden = True

    dv_staff = DataValidation(type="list", formula1=f"={STAFF_NAMES}", allow_blank=True)
    dv_site = DataValidation(type="list", formula1=f"={SITE_NAMES}", allow_blank=True)
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
    dv_time = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988", allow_blank=True)
    dv_break = DataValidation(type="whole", operator="between", formula1="0", formula2="240", allow_blank=True)
    for dv, cols in ((dv_staff, "B"), (dv_site, "C"), (dv_date, "A"), (dv_time, "DE"), (dv_break, "F")):
        dv.error, dv.showErrorMessage = "Pick from the list or enter a valid value.", True
        ws.add_data_validation(dv)
        for col in cols:
            dv.add(f"{col}{LOG_FIRST}:{col}{LOG_LAST}")
    ws.conditional_formatting.add(
        f"A{LOG_FIRST}:K{LOG_LAST}",
        FormulaRule(formula=[f'AND($K{LOG_FIRST}<>"",$K{LOG_FIRST}<>"OK")'], fill=FLAG_FILL),
    )


def build_staff(ws):
    header(ws, 1, ["Employee", "Region", "Role", "Hourly rate", "Active (Y/N)", "Notes"], [24, 20, 16, 12, 12, 40])
    ws.freeze_panes = "A2"
    for r, row in enumerate(PLACEHOLDER_STAFF, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=4).number_format = '"$"#,##0.00'
    dv = DataValidation(type="list", formula1='"Yellowknife,Richmond/Vancouver"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{STAFF_LAST}")
    dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"E2:E{STAFF_LAST}")
    ws["H1"] = "Names must be unique. Set Active to N instead of deleting someone who has logged hours."


def build_sites(ws):
    data = json.loads((HERE / "sites.json").read_text(encoding="utf-8"))
    header(ws, 1, ["Client / Site", "Region", "Type", "Active (Y/N)"], [44, 20, 14, 12])
    ws.freeze_panes = "A2"
    for r, s in enumerate(data["sites"], start=2):
        ws.cell(row=r, column=1, value=s["site"])
        ws.cell(row=r, column=2, value=s["region"])
        ws.cell(row=r, column=3, value=s["type"])
        ws.cell(row=r, column=4, value=s["active"])
    dv = DataValidation(type="list", formula1='"Yellowknife,Richmond/Vancouver"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{SITES_LAST}")
    ws["F1"] = "Use the client name exactly as it appears in Jobber. Source: Jobber invoices 2026-08-12 to 2026-10-09 + NACC."


def build_pay_period(ws):
    ws["A1"], ws["A1"].font = "Pay Period", TITLE_FONT
    # Pay is monthly: the period is a calendar month.
    ws["A2"], ws["B2"] = "Months back (0 = this month, 1 = last month)", 0
    ws["B2"].fill = INPUT_FILL
    ws["A3"] = "Pay is monthly. Run payroll from last month (Months back = 1)."
    ws["A5"], ws["B5"] = "Period start", "=DATE(YEAR(TODAY()),MONTH(TODAY())-B2,1)"
    ws["A6"], ws["B6"] = "Period end", "=DATE(YEAR(TODAY()),MONTH(TODAY())-B2+1,0)"
    for c in ("B5", "B6"):
        ws[c].number_format = DATE_FMT
        ws[c].font = BOLD
    header_row = 8
    labels = ["Employee", "Hours", "Shifts", "Hourly rate", "Gross pay", "Rows flagged"]
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for col, w in zip("ABCDEF", (46, 10, 10, 12, 12, 13)):
        ws.column_dimensions[col].width = w
    dates = f'{RNG("A")},">="&$B$5,{RNG("A")},"<="&$B$6'
    first = header_row + 1
    last = first + (STAFF_LAST - 2)
    for i, r in enumerate(range(first, last + 1)):
        ws[f"A{r}"] = f'=IF(Staff!A{2 + i}="","",Staff!A{2 + i})'
        ws[f"B{r}"] = f'=IF(A{r}="","",SUMIFS({RNG("G")},{RNG("B")},A{r},{dates}))'
        ws[f"C{r}"] = f'=IF(A{r}="","",COUNTIFS({RNG("B")},A{r},{dates}))'
        ws[f"D{r}"] = f'=IF(A{r}="","",IFERROR(INDEX({STAFF_RATES},MATCH(A{r},{STAFF_NAMES},0)),0))'
        ws[f"E{r}"] = f'=IF(A{r}="","",B{r}*D{r})'
        ws[f"F{r}"] = f'=IF(A{r}="","",COUNTIFS({RNG("B")},A{r},{dates},{RNG("K")},"<>OK"))'
        ws[f"B{r}"].number_format = "0.00"
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = '"$"#,##0.00'
    tot = last + 1
    ws[f"A{tot}"], ws[f"A{tot}"].font = "TOTAL", BOLD
    for col in "BCEF":
        ws[f"{col}{tot}"] = f"=SUM({col}{first}:{col}{last})"
        ws[f"{col}{tot}"].font = BOLD
    ws[f"B{tot}"].number_format = "0.00"
    ws[f"E{tot}"].number_format = '"$"#,##0.00'
    ws.conditional_formatting.add(f"F{first}:F{last}", FormulaRule(formula=[f"N(F{first})>0"], fill=FLAG_FILL))
    ws.freeze_panes = f"A{first}"
    ws["H2"] = "Payroll steps: 1) Rows flagged = 0 for everyone.  2) Hours match Jobber timesheets (Weekly tab).  3) Run payroll from this table."


def build_weekly(ws):
    ws["A1"], ws["A1"].font = "Weekly", TITLE_FONT
    ws["A2"], ws["B2"] = "Weeks back (0 = this week)", 0
    ws["B2"].fill = INPUT_FILL
    ws["A3"], ws["B3"] = "Week of (Mon)", "=TODAY()-WEEKDAY(TODAY(),3)-7*B2"
    ws["B3"].number_format, ws["B3"].font = DATE_FMT, BOLD
    ws.column_dimensions["A"].width = 30

    # Section 1: hours per employee per day
    hr = 5
    ws.cell(row=hr, column=1, value="Employee")
    for d in range(7):
        c = ws.cell(row=hr, column=2 + d, value=f"=$B$3+{d}")
        c.number_format = "ddd mm-dd"
    ws.cell(row=hr, column=9, value="Total")
    for col in range(1, 10):
        cell = ws.cell(row=hr, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    first = hr + 1
    last = first + (STAFF_LAST - 2)
    for i, r in enumerate(range(first, last + 1)):
        ws[f"A{r}"] = f'=IF(Staff!A{2 + i}="","",Staff!A{2 + i})'
        for d in range(7):
            col = ws.cell(row=hr, column=2 + d).column_letter
            ws[f"{col}{r}"] = f'=IF($A{r}="","",SUMIFS({RNG("G")},{RNG("B")},$A{r},{RNG("A")},{col}${hr}))'
            ws[f"{col}{r}"].number_format = "0.00;-0.00;"
        ws[f"I{r}"] = f'=IF($A{r}="","",SUM(B{r}:H{r}))'
        ws[f"I{r}"].number_format, ws[f"I{r}"].font = "0.00", BOLD

    # Section 2: site cross-check against Jobber (columns K:Q)
    labels = ["Client / Site", "Shifts logged", "Crew hours logged", "Jobber visits (enter)",
              "Jobber hours (enter)", "Difference", "Flag"]
    for i, label in enumerate(labels):
        c = ws.cell(row=hr, column=11 + i, value=label)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True)
    for col, w in zip("KLMNOPQ", (40, 9, 10, 10, 10, 10, 28)):
        ws.column_dimensions[col].width = w
    ws["K4"] = "Cross-check: enter last week's Jobber visit count and timesheet hours per site in N and O. Clear N:O when the week changes."
    site_first, site_last = first, first + (SITES_LAST - 2)
    for i, r in enumerate(range(site_first, site_last + 1)):
        ws[f"K{r}"] = f'=IF(Sites!A{2 + i}="","",Sites!A{2 + i})'
        ws[f"L{r}"] = f'=IF(K{r}="","",COUNTIFS({RNG("C")},K{r},{RNG("I")},$B$3))'
        ws[f"M{r}"] = f'=IF(K{r}="","",SUMIFS({RNG("G")},{RNG("C")},K{r},{RNG("I")},$B$3))'
        ws[f"N{r}"].fill = ws[f"O{r}"].fill = INPUT_FILL
        ws[f"P{r}"] = f'=IF(OR(K{r}="",O{r}=""),"",M{r}-O{r})'
        ws[f"Q{r}"] = (
            f'=IF(K{r}="","",'
            f'IF(AND(N(N{r})>0,L{r}=0),"Jobber visit, no hours logged",'
            f'IF(AND(L{r}>0,N{r}=""),"Enter Jobber visits",'
            f'IF(AND(L{r}>0,N(N{r})=0),"Hours logged, no Jobber visit",'
            f'IF(AND(O{r}<>"",ABS(N(P{r}))>0.5),"Hours differ from Jobber","")))))'
        )
        ws[f"M{r}"].number_format = ws[f"P{r}"].number_format = "0.00;-0.00;"
    ws.conditional_formatting.add(f"Q{site_first}:Q{site_last}",
                                  FormulaRule(formula=[f'Q{site_first}<>""'], fill=FLAG_FILL))
    ws.freeze_panes = f"B{first}"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Start Here"
    build_start_here(ws)
    build_hours_log(wb.create_sheet("Hours Log"))
    build_staff(wb.create_sheet("Staff"))
    build_sites(wb.create_sheet("Sites"))
    build_pay_period(wb.create_sheet("Pay Period"))
    build_weekly(wb.create_sheet("Weekly"))
    for name, color in (("Start Here", NAVY), ("Hours Log", "2E7D32"), ("Pay Period", "B26A00"), ("Weekly", "B26A00")):
        wb[name].sheet_properties.tabColor = color
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
